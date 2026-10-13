"""Pricing worksheet workbook (spec 12.3): one row per RFQ line, filled by band."""

from __future__ import annotations

from collections.abc import Callable, Sequence
from typing import get_args

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from rfq_agent.schemas import (
    Band,
    CoveragePath,
    ResolvedRow,
    RunRecord,
    RunSummary,
    WorksheetRow,
)

# Excel's built-in good / neutral / bad fills.
BAND_COLOURS: dict[Band, str] = {"green": "C6EFCE", "yellow": "FFEB9C", "red": "FFC7CE"}
BAND_FILLS: dict[Band, PatternFill] = {
    band: PatternFill(fill_type="solid", start_color=rgb, end_color=rgb)
    for band, rgb in BAND_COLOURS.items()
}
REASONING_PATHS: tuple[CoveragePath, ...] = get_args(CoveragePath)  # the four agent paths
ELAPSED_NOTE = "elapsed_s: wall clock from run creation to finalize, including any downtime"

Cell = str | int | float | bool | None
Column = tuple[str, Callable[[WorksheetRow], Cell]]


def build_workbook(record: RunRecord, rows: Sequence[ResolvedRow], summary: RunSummary) -> Workbook:
    """The worksheet sheet (rows as shown, review overlays included) and the run summary."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "worksheet"
    _write_rows(sheet, [r.worksheet for r in rows])
    _write_summary(book.create_sheet("summary"), record, summary)
    return book


def columns(sites: Sequence[str]) -> list[Column]:
    """Worksheet columns in order: identity, OE, the price reference side by side, details."""

    def site_price(site: str) -> Callable[[WorksheetRow], Cell]:
        return lambda w: w.retail_prices.get(site)

    def reasoning(path: str) -> Callable[[WorksheetRow], Cell]:
        return lambda w: w.reasoning_summary.get(path)

    return [
        ("row_id", lambda w: w.row_id),
        ("customer_part_no_raw", lambda w: w.customer_part_no_raw),
        ("customer_part_no", lambda w: w.customer_part_no),
        ("our_part_no", lambda w: w.our_part_no),
        ("oe_status", lambda w: w.oe_status),
        ("current_oe", lambda w: w.current_oe),
        ("supersession_chain", lambda w: " -> ".join(w.supersession_chain)),
        ("oe_candidates", lambda w: ", ".join(w.oe_candidates)),
        ("msrp", lambda w: w.msrp),
        ("selected_retail_price", lambda w: w.selected_retail_price),
        ("landed_cost", lambda w: w.landed_cost),
        ("msrp_status", lambda w: w.msrp_status),
        *[(f"retail_price_{site}", site_price(site)) for site in sites],
        ("selected_retail_source", lambda w: w.selected_retail_source),
        ("retail_outliers", lambda w: ", ".join(w.retail_outliers)),
        ("should_cost_fob", lambda w: w.should_cost_fob),
        ("shared_oe_rows", lambda w: ", ".join(str(r) for r in w.shared_oe_rows)),
        ("price_source_row", lambda w: w.price_source_row),
        ("shared_prices_disagree", lambda w: w.shared_prices_disagree),
        ("confidence_score", lambda w: w.confidence_score),
        ("band", lambda w: w.band),
        ("review_status", lambda w: w.review_status),
        ("review_reasons", lambda w: "\n".join(w.review_reasons)),
        ("failures", lambda w: "\n".join(w.failures)),
        *[(f"reasoning_{path}", reasoning(path)) for path in REASONING_PATHS],
    ]


def _write_rows(sheet: Worksheet, rows: Sequence[WorksheetRow]) -> None:
    sites = sorted({site for w in rows for site in w.retail_prices})
    cols = columns(sites)
    _header(sheet, [name for name, _ in cols])
    for w in rows:
        sheet.append([value(w) for _, value in cols])
        if w.band is not None:
            for cell in sheet[sheet.max_row]:
                cell.fill = BAND_FILLS[w.band]
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions


def _write_summary(sheet: Worksheet, record: RunRecord, summary: RunSummary) -> None:
    spec = record.spec
    items: list[tuple[str, Cell]] = [
        ("run_id", record.run_id),
        ("status", record.status),
        ("input", spec.input),
        ("strategy", spec.strategy),
        ("batch_size", spec.batch_size),
        ("review", spec.review),
        ("created_at", record.created_at),
        ("finished_at", record.finished_at),
        ("rows", summary.rows),
        *[(f"rows_{band}", n) for band, n in summary.bands.items()],
        *[(f"reviewed_{status}", n) for status, n in summary.reviews.items()],
        ("review_pending", summary.review_pending),
        *[(f"coverage_{path}", share) for path, share in summary.coverage.items()],
        ("rows_with_failure", summary.rows_with_failure),
        ("llm_calls", summary.llm_calls),
        ("input_tokens", summary.input_tokens),
        ("output_tokens", summary.output_tokens),
        ("llm_cost_usd", summary.llm_cost_usd),
        ("elapsed_s", summary.elapsed_s),
    ]
    _header(sheet, ["metric", "value"])
    for item in items:
        sheet.append(list(item))
    sheet.append([])
    sheet.append([ELAPSED_NOTE])


def _header(sheet: Worksheet, names: list[str]) -> None:
    sheet.append(names)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
