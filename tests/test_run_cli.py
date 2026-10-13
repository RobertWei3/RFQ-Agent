"""``rfq run`` / ``rfq rerun-failed`` / ``rfq runs show`` (spec 13), fake mode, real MCP."""

from __future__ import annotations

import os
import re
import signal
import subprocess
import sys
import time
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from typer.testing import CliRunner

from rfq_agent.cli import app
from rfq_agent.config import load_settings
from rfq_agent.schemas import ResolvedRow
from rfq_agent.store.runs import RunStore


@pytest.fixture
def cli_config(config_dir: Path, generated_dataset: Path) -> Path:
    path = config_dir / "settings.yaml"
    data = yaml.safe_load(path.read_text())
    data["data"]["dataset_dir"] = str(generated_dataset)
    path.write_text(yaml.safe_dump(data))
    return config_dir


def _run_id(output: str) -> str:
    found = re.search(r"run_id=(\S+)", output)
    assert found, output
    return found.group(1)


def _sample(generated_dataset: Path) -> str:
    return str(generated_dataset / "rfq_sample_20.xlsx")


def _store(config_dir: Path) -> RunStore:
    return RunStore.open(load_settings(config_dir).db_path)


def test_run_processes_the_rfq_in_batches_and_exports_the_rows(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    result = CliRunner().invoke(
        app,
        [
            "run",
            "--input",
            _sample(generated_dataset),
            "--batch-size",
            "7",
            "--output-dir",
            str(tmp_path),
        ],
    )

    assert result.exit_code == 0, result.output
    run_id = _run_id(result.output)
    exported = tmp_path / f"run_{run_id}.jsonl"
    rows = [ResolvedRow.model_validate_json(x) for x in exported.read_text().splitlines()]
    assert [r.line.row_id for r in rows] == list(range(1, 21))
    assert all(r.confidence is not None for r in rows)
    assert exported.with_suffix(".compare.md").is_file()
    assert exported.with_suffix(".trace.jsonl").is_file()
    book = load_workbook(exported.with_suffix(".xlsx"))
    assert [s.title for s in book.worksheets] == ["worksheet", "summary"]
    assert book.worksheets[0].max_row == 21
    assert "batches: completed 3" in result.output
    assert "LLM calls: oe_interchange" in result.output

    shown = CliRunner().invoke(app, ["runs", "show", run_id])
    assert shown.exit_code == 0, shown.output
    assert "status completed" in shown.output
    assert "batch 3  rows 15-20  completed  attempts 1" in shown.output


def test_resume_and_rerun_failed_on_a_finished_run_change_nothing(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    runner = CliRunner()
    first = runner.invoke(
        app, ["run", "--input", _sample(generated_dataset), "--output-dir", str(tmp_path)]
    )
    run_id = _run_id(first.output)

    resumed = runner.invoke(app, ["run", "--resume", run_id, "--output-dir", str(tmp_path)])
    rerun = runner.invoke(app, ["rerun-failed", run_id, "--output-dir", str(tmp_path)])

    assert resumed.exit_code == 0, resumed.output
    assert rerun.exit_code == 0, rerun.output
    assert "nothing to rerun" in rerun.output
    store = _store(cli_config)
    assert [b.attempts for b in store.batches(run_id)] == [1]
    store.close()


def test_real_provider_needs_allow_real_and_creates_no_run(
    cli_config: Path, generated_dataset: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("RFQ_LLM_PROVIDER")

    result = CliRunner().invoke(app, ["run", "--input", _sample(generated_dataset)])

    assert result.exit_code == 2
    assert "--allow-real" in result.output
    assert "run_id=" not in result.output


def test_review_prompts_red_rows_and_quitting_leaves_the_run_resumable(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    runner = CliRunner()
    started = runner.invoke(
        app,
        ["run", "--input", _sample(generated_dataset), "--review", "--output-dir", str(tmp_path)],
        input="a\nq\n",
    )

    assert started.exit_code == 0, started.output
    run_id = _run_id(started.output)
    assert "[a]ccept" in started.output
    assert "review paused" in started.output
    assert (tmp_path / f"run_{run_id}.jsonl").is_file()  # finalized before review

    resumed = runner.invoke(
        app,
        ["run", "--resume", run_id, "--output-dir", str(tmp_path)],
        input="r\n\np\n42\n\n" + "a\n" * 20,  # reject (empty note), then MSRP 42
    )

    assert resumed.exit_code == 0, resumed.output
    assert (tmp_path / f"run_{run_id}.jsonl").is_file()
    store = _store(cli_config)
    log = store.reviews(run_id)
    status = store.run(run_id).status
    store.close()
    assert status == "completed"
    assert [e.action for e in log[:3]] == ["accept", "reject", "modify_price"]
    assert (log[2].msrp_after, log[2].retail_after) == (42.0, None)


def _completed_batches(config_dir: Path, run_id: str) -> int:
    store = _store(config_dir)
    try:
        return sum(b.status == "completed" for b in store.batches(run_id))
    finally:
        store.close()


def test_a_killed_run_resumes_without_redoing_completed_batches(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    """Spec 13 acceptance, small: SIGKILL mid-run, --resume, same rows as an unbroken run."""
    args = ["run", "--input", str(generated_dataset / "rfq_200.xlsx"), "--batch-size", "20"]
    args += ["--output-dir", str(tmp_path)]
    unbroken = _run_id(CliRunner().invoke(app, args).output)

    proc = subprocess.Popen(
        [sys.executable, "-m", "rfq_agent.cli", *args],
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        env=os.environ.copy(),
    )
    assert proc.stdout is not None
    run_id = ""
    while not run_id:
        line = proc.stdout.readline()
        assert line, "run exited before printing its run_id"
        run_id = _run_id(line) if "run_id=" in line else ""
    deadline = time.monotonic() + 60
    while _completed_batches(cli_config, run_id) < 3:
        assert proc.poll() is None, "run finished before it could be killed"
        assert time.monotonic() < deadline
        time.sleep(0.02)
    proc.send_signal(signal.SIGKILL)
    proc.wait()
    store = _store(cli_config)
    done_before = [b.batch_id for b in store.batches(run_id) if b.status == "completed"]
    store.close()
    assert 3 <= len(done_before) < 10

    resumed = CliRunner().invoke(app, ["run", "--resume", run_id, "--output-dir", str(tmp_path)])

    assert resumed.exit_code == 0, resumed.output
    store = _store(cli_config)
    batches = store.batches(run_id)
    store.close()
    assert all(b.status == "completed" for b in batches)
    assert all(b.attempts == 1 for b in batches if b.batch_id in done_before)
    exported = (tmp_path / f"run_{run_id}.jsonl").read_text()
    assert exported == (tmp_path / f"run_{unbroken}.jsonl").read_text()


def test_resume_announces_only_the_rows_a_pending_rerun_will_run(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    runner = CliRunner()
    first = runner.invoke(
        app,
        ["run", "--input", _sample(generated_dataset), "--batch-size", "7"]
        + ["--output-dir", str(tmp_path)],
    )
    run_id = _run_id(first.output)
    store = _store(cli_config)
    store.reopen_for_rerun(run_id, {2: [9], 3: [16, 17]})  # as a killed rerun leaves it
    store.close()

    resumed = runner.invoke(app, ["run", "--resume", run_id, "--output-dir", str(tmp_path)])

    assert resumed.exit_code == 0, resumed.output
    assert "rows=3 " in resumed.output
    assert "oe_interchange (fake/fake) expected ~9," in resumed.output


def test_resume_is_refused_while_another_process_drives_the_run(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    runner = CliRunner()
    args = ["--output-dir", str(tmp_path)]
    run_id = _run_id(
        runner.invoke(app, ["run", "--input", _sample(generated_dataset), *args]).output
    )
    store = _store(cli_config)
    store.reopen_for_rerun(run_id, {1: [1]})
    store.acquire_lease(run_id, "api-worker", ttl_s=60)
    store.close()

    refused = runner.invoke(app, ["run", "--resume", run_id, *args])

    assert refused.exit_code == 3
    assert f"run {run_id} is driven by api-worker" in refused.output

    handed_over = runner.invoke(
        app, ["run", "--resume", run_id, "--lease-token", "api-worker", *args]
    )
    assert handed_over.exit_code == 0, handed_over.output
    assert "batches: completed 1" in handed_over.output


def test_the_worksheet_is_exported_before_the_first_review_question(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    # No answer at all (like Ctrl-C at the first prompt): the run is finalized and the
    # reviewer already has the files (spec 13: review happens on a finalized run).
    aborted = CliRunner().invoke(
        app,
        ["run", "--input", _sample(generated_dataset), "--review", "--output-dir", str(tmp_path)],
        input="",
    )

    assert aborted.exit_code != 0
    run_id = _run_id(aborted.output)
    assert (tmp_path / f"run_{run_id}.jsonl").is_file()
    assert (tmp_path / f"run_{run_id}.xlsx").is_file()
    store = _store(cli_config)
    status = store.run(run_id).status
    store.close()
    assert status == "awaiting_review"


def test_a_worker_that_fails_before_taking_over_releases_the_lease_it_was_given(
    cli_config: Path, generated_dataset: Path, tmp_path: Path
) -> None:
    runner = CliRunner()
    run_id = _run_id(
        runner.invoke(
            app, ["run", "--input", _sample(generated_dataset), "--output-dir", str(tmp_path)]
        ).output
    )
    store = _store(cli_config)
    store.acquire_lease(run_id, "api-token", ttl_s=60)

    failed = runner.invoke(
        app, ["runs", "answer", run_id, "--decision", "{not json", "--lease-token", "api-token"]
    )

    assert failed.exit_code != 0
    assert store.lease(run_id) is None
    store.close()
