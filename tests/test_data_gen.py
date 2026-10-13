"""Generator acceptance tests (spec section 17, P1)."""

from __future__ import annotations

import hashlib
import json
import sqlite3
import statistics
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rfq_agent.config import DEFAULT_CONFIG_DIR, PROJECT_ROOT, load_settings
from rfq_agent.data_gen.audit import load_pages, load_truth, measure, parse_spec_page
from rfq_agent.data_gen.params import SECTION_6_3, Difficulty, GenParams
from rfq_agent.data_gen.truth import GroundTruthRow
from rfq_agent.data_gen.world import RETAIL_SITES, expected_selected_site
from rfq_agent.data_gen.writer import generate_dataset
from rfq_agent.partno import normalize_part_no

N = 200


@pytest.fixture(scope="module")
def dataset(generated_dataset: Path) -> Path:
    return generated_dataset


@pytest.fixture(scope="module")
def truth(dataset: Path) -> list[GroundTruthRow]:
    return load_truth(dataset)


def _digest(directory: Path) -> dict[str, str]:
    result = {}
    for path in sorted(directory.rglob("*")):
        if path.suffix in {".jsonl", ".json"}:
            result[str(path.relative_to(directory))] = hashlib.sha256(path.read_bytes()).hexdigest()
    conn = sqlite3.connect(directory / "internal.db")
    result["internal.db"] = hashlib.sha256("\n".join(conn.iterdump()).encode()).hexdigest()
    conn.close()
    wb = load_workbook(directory / f"rfq_{N}.xlsx", read_only=True)
    ws = wb.active
    assert ws is not None
    cells = repr(list(ws.iter_rows(values_only=True)))
    wb.close()
    result["xlsx"] = hashlib.sha256(cells.encode()).hexdigest()
    return result


def test_ratios_measured_from_artifacts_within_3pct(dataset: Path) -> None:
    measured = measure(dataset)
    targets = GenParams().ratios
    for difficulty in SECTION_6_3:
        realized = measured[difficulty] / N
        assert abs(realized - targets[difficulty]) <= 0.03, (difficulty, realized)


def test_deterministic(dataset: Path, tmp_path: Path) -> None:
    cost_model = load_settings(DEFAULT_CONFIG_DIR).cost_model
    generate_dataset(GenParams(n_parts=N, seed=42), cost_model, tmp_path)
    assert _digest(tmp_path) == _digest(dataset)


def test_different_seed_differs(dataset: Path, tmp_path: Path) -> None:
    cost_model = load_settings(DEFAULT_CONFIG_DIR).cost_model
    generate_dataset(GenParams(n_parts=N, seed=7), cost_model, tmp_path)
    assert _digest(tmp_path)["ground_truth.jsonl"] != _digest(dataset)["ground_truth.jsonl"]


def test_part_numbers_unique_after_normalization(truth: list[GroundTruthRow]) -> None:
    customer = [normalize_part_no(t.customer_part_no) for t in truth]
    assert len(customer) == len(set(customer))
    oes = {oe for t in truth for oe in [*t.accepted_oes, *t.supersession_chain]}
    assert len({normalize_part_no(o) for o in oes}) == len(oes)
    assert not set(customer) & {normalize_part_no(o) for o in oes}


def test_supersession_chains(truth: list[GroundTruthRow], dataset: Path) -> None:
    catalog = load_pages(dataset, "oem_catalog")
    for t in truth:
        chain = t.supersession_chain
        assert len(chain) == len(set(chain))  # acyclic
        if chain:
            assert chain[-1] == t.primary_oe
        for old, new in zip(chain, chain[1:], strict=False):
            assert new in catalog[f"cat-{normalize_part_no(old)}"]["text"]
            assert old in catalog[f"cat-{normalize_part_no(new)}"]["text"]


def test_oe_status_and_accepted_set(truth: list[GroundTruthRow]) -> None:
    for t in truth:
        if t.oe_status_expected == "not_found":
            assert t.primary_oe is None and t.accepted_oes == []
        else:
            assert t.primary_oe == t.accepted_oes[0]
            assert (t.oe_status_expected == "multiple") == (len(t.accepted_oes) > 1)


def test_many_to_one_rows_share_oe(truth: list[GroundTruthRow]) -> None:
    by_id = {t.row_id: t for t in truth}
    for t in truth:
        for other in t.shares_oe_with:
            assert by_id[other].primary_oe == t.primary_oe is not None


def test_retail_outliers_and_selection(truth: list[GroundTruthRow]) -> None:
    for t in truth:
        listings = {s: lst for s, lst in t.retail.sites.items() if lst is not None}
        if not listings:
            continue
        median = statistics.median(lst.effective_price for lst in listings.values())
        for site, lst in listings.items():
            deviation = abs(lst.effective_price - median) / median
            if lst.outlier:
                assert deviation > 0.40, (t.row_id, site)
            else:
                assert deviation <= 0.40, (t.row_id, site)
        assert t.retail.expected_selected_site == expected_selected_site(t.retail.sites)
        assert t.retail.expected_selected_site not in t.retail.outlier_sites


def test_price_levels(truth: list[GroundTruthRow]) -> None:
    ranges = GenParams().msrp_ranges
    for t in truth:
        if t.msrp.value is not None:
            r = ranges[t.spec.category]
            assert r.low <= t.msrp.value <= r.high
        lists = [lst.list_price for lst in t.retail.sites.values() if lst and not lst.outlier]
        if lists and t.msrp.value is not None:
            share = t.cost.landed_cost / statistics.median(lists)
            assert 0.20 <= share <= 0.55, (t.row_id, share)
            assert all(0.549 * t.msrp.value <= p <= 0.951 * t.msrp.value for p in lists)


def test_msrp_cases(truth: list[GroundTruthRow], dataset: Path) -> None:
    msrp_pages = load_pages(dataset, "msrp")
    for t in truth:
        if t.primary_oe is None:
            continue
        current_page = msrp_pages.get(f"msrp-{normalize_part_no(t.primary_oe)}")
        if "msrp_fallback" in t.tags:
            assert t.msrp.source_oe == t.supersession_chain[-2]
            assert current_page is None
        if t.msrp.status == "not_found":
            assert t.msrp.value is None and current_page is None
        else:
            assert t.msrp.source_oe is not None
            page = msrp_pages[f"msrp-{normalize_part_no(t.msrp.source_oe)}"]
            assert f"{t.msrp.value:.2f}" in page["text"].replace(",", "")
            assert ("discontinued" in page["text"].lower()) == (t.msrp.status == "discontinued")


def test_external_spec_pages_match_truth_unless_conflict(
    truth: list[GroundTruthRow], dataset: Path
) -> None:
    pages = load_pages(dataset, "external_spec")
    for t in truth:
        page = pages.get(f"spec-{normalize_part_no(t.primary_oe or '')}")
        if page is None:
            assert not {"internal_spec_missing", "spec_conflict"} & set(t.tags)
            continue
        material, weight = parse_spec_page(page["text"])
        agrees = (
            material == t.spec.material and abs(weight - t.spec.weight_kg) / t.spec.weight_kg < 0.02
        )  # type: ignore[operator]
        assert agrees != ("spec_conflict" in t.tags), t.row_id


def test_internal_db_contents(truth: list[GroundTruthRow], dataset: Path) -> None:
    conn = sqlite3.connect(dataset / "internal.db")
    xref = dict(conn.execute("SELECT customer_part_no_norm, our_part_no FROM customer_xref"))
    parts = {
        r[0]: r[1:] for r in conn.execute("SELECT our_part_no, material, weight_kg FROM our_parts")
    }
    prices = dict(conn.execute("SELECT material, price_per_kg FROM material_prices"))
    conn.close()
    assert len(parts) > len({t.our_part_no for t in truth})  # distractors present
    for t in truth:
        key = normalize_part_no(t.customer_part_no)
        assert (key in xref) == (t.our_part_reachability == "xref")
        if key in xref:
            assert xref[key] == t.our_part_no
        assert parts[t.our_part_no] == (t.internal_spec.material, t.internal_spec.weight_kg)
        assert t.spec.material in prices


def test_rfq_workbook(truth: list[GroundTruthRow], dataset: Path) -> None:
    wb = load_workbook(dataset / f"rfq_{N}.xlsx", read_only=True)
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    assert rows[0] == (
        "Customer Part No",
        "Our Part No",
        "Description",
        "Annual Qty",
        "Customer Name",
    )
    data = [r for r in rows[1:] if any(v is not None for v in r)]
    assert len(rows) - 1 > len(data)  # blank rows present
    assert len(data) == N
    for t, row in zip(truth, data, strict=True):
        assert row[0] == t.customer_part_no_raw
        assert normalize_part_no(str(row[0])) == normalize_part_no(t.customer_part_no)
        assert (row[1] is not None) == (t.our_part_reachability == "given")


def test_sample_covers_every_difficulty(dataset: Path) -> None:
    with (dataset / "rfq_sample_20_ground_truth.jsonl").open() as fh:
        sample = [GroundTruthRow.model_validate_json(line) for line in fh]
    assert len(sample) == 20
    assert [s.row_id for s in sample] == list(range(1, 21))
    tags = {tag for s in sample for tag in s.tags}
    assert tags == {d.value for d in Difficulty}
    sources = {s.source_row_id for s in sample}
    for s in sample:  # many-to-one partners are in the sample too
        assert set(s.shares_oe_with) <= sources
    assert (dataset / "sample_copy.xlsx").is_file()


def test_manifest(dataset: Path) -> None:
    manifest = json.loads((dataset / "manifest.json").read_text())
    assert manifest["rows"] == N
    assert set(manifest["corpus_sizes"]) == {
        "oem_catalog",
        "interchange",
        "msrp",
        "external_spec",
        *RETAIL_SITES,
    }


def test_insufficient_pool_raises() -> None:
    cost_model = load_settings(DEFAULT_CONFIG_DIR).cost_model
    ratios = GenParams().ratios
    ratios[Difficulty.SPEC_CONFLICT] = 0.9
    with pytest.raises(ValueError, match="spec_conflict"):
        from rfq_agent.data_gen.world import build_world

        build_world(GenParams(n_parts=50, ratios=ratios), cost_model)


def test_regenerating_the_sample_leaves_the_committed_workbook_unchanged(dataset: Path) -> None:
    # `rfq data generate` overwrites data/samples; a fresh clone must stay clean.
    committed = PROJECT_ROOT / "data" / "samples" / "rfq_sample_20.xlsx"
    assert (dataset / "sample_copy.xlsx").read_bytes() == committed.read_bytes()
