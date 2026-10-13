"""Pricing worksheet workbook (spec 12.3; phase-8 Q12-Q15)."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from rfq_agent.output.excel import build_workbook
from rfq_agent.schemas import ResolvedRow, RunRecord, RunSpec, RunSummary
from tests.builders import msrp, oe, resolved_row, retail

GREEN, YELLOW, RED = "C6EFCE", "FFEB9C", "FFC7CE"


def _record() -> RunRecord:
    spec = RunSpec(
        input="rfq.xlsx", dataset_dir="x", batch_size=2, strategy="evidence", review=True
    )
    return RunRecord(
        run_id="run-1",
        created_at="2026-10-08T10:00:00+00:00",
        spec=spec,
        status="awaiting_review",
        finished_at="2026-10-08T10:01:30+00:00",
    )


def _summary() -> RunSummary:
    return RunSummary(
        rows=4,
        bands={"green": 1, "yellow": 1, "red": 2},
        reviews={"accepted": 0, "modified": 0, "rejected": 1},
        review_pending=1,
        coverage={"oe": 0.75, "msrp": 0.5, "retail": 0.75, "should_cost": 1.0},
        rows_with_failure=0,
        llm_calls=12,
        input_tokens=3400,
        output_tokens=560,
        llm_cost_usd=0.0021,
        elapsed_s=90.0,
    )


def _rows() -> list[ResolvedRow]:
    chained = oe(2, "NEW-3", supersession_chain=["OLD-1", "MID-2", "NEW-3"])
    rejected = resolved_row(3, "red", retail=retail(3, {"retail_2": 70.0, "retail_5": 75.0}))
    sheet = rejected.worksheet.model_copy(
        update={"review_status": "rejected", "review_reasons": ["no MSRP", "price conflict"]}
    )
    unscored = resolved_row(4, "red", msrp=msrp(4, None))
    return [
        resolved_row(1, "green"),
        resolved_row(2, "yellow", oe=chained),
        rejected.model_copy(update={"worksheet": sheet}),
        unscored.model_copy(
            update={"worksheet": unscored.worksheet.model_copy(update={"band": None})}
        ),
    ]


def _reopen(book: Workbook) -> Workbook:
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return load_workbook(buffer)


def _sheet() -> Worksheet:
    book = _reopen(build_workbook(_record(), _rows(), _summary()))
    sheet = book.worksheets[0]
    assert sheet.title == "worksheet"
    return sheet


def _header(sheet: Worksheet) -> list[str]:
    return [str(c.value) for c in sheet[1]]


def _by_header(sheet: Worksheet, row: int) -> dict[str, object]:
    return dict(zip(_header(sheet), [c.value for c in sheet[row]], strict=True))


def test_columns_put_the_price_reference_side_by_side_and_one_column_per_site() -> None:
    assert _header(_sheet()) == [
        "row_id",
        "customer_part_no_raw",
        "customer_part_no",
        "our_part_no",
        "oe_status",
        "current_oe",
        "supersession_chain",
        "oe_candidates",
        "msrp",
        "selected_retail_price",
        "landed_cost",
        "msrp_status",
        "retail_price_retail_1",
        "retail_price_retail_2",
        "retail_price_retail_3",
        "retail_price_retail_5",
        "selected_retail_source",
        "retail_outliers",
        "should_cost_fob",
        "shared_oe_rows",
        "price_source_row",
        "shared_prices_disagree",
        "confidence_score",
        "band",
        "review_status",
        "review_reasons",
        "failures",
        "reasoning_oe",
        "reasoning_msrp",
        "reasoning_retail",
        "reasoning_should_cost",
    ]


def test_cells_show_the_worksheet_values() -> None:
    sheet = _sheet()
    second = _by_header(sheet, 3)
    third = _by_header(sheet, 4)

    assert second["supersession_chain"] == "OLD-1 -> MID-2 -> NEW-3"
    assert (second["msrp"], second["selected_retail_price"], second["landed_cost"]) == (
        100.0,
        80.0,
        30.0,
    )
    assert second["retail_price_retail_3"] == 85.0
    assert second["retail_price_retail_5"] is None
    assert second["reasoning_msrp"] == "msrp reasoning 2"
    assert third["review_status"] == "rejected"
    assert third["review_reasons"] == "no MSRP\nprice conflict"
    assert third["retail_price_retail_1"] is None


def test_each_row_is_filled_with_its_band_colour_and_the_header_is_frozen() -> None:
    sheet = _sheet()
    width = len(_header(sheet))

    def fills(row: int) -> set[str | None]:
        cells = sheet[row]
        assert len(cells) == width
        return {c.fill.fgColor.rgb[-6:] if c.fill.fill_type == "solid" else None for c in cells}

    assert [fills(r) for r in (2, 3, 4, 5)] == [{GREEN}, {YELLOW}, {RED}, {None}]
    assert sheet.freeze_panes == "B2"
    assert sheet.auto_filter.ref == f"A1:{sheet.cell(1, width).column_letter}5"


def test_the_second_sheet_summarizes_the_run() -> None:
    book = _reopen(build_workbook(_record(), _rows(), _summary()))
    summary = book.worksheets[1]
    values = {row[0]: row[1] for row in summary.iter_rows(values_only=True) if row[0]}

    assert summary.title == "summary"
    assert values["run_id"] == "run-1"
    assert values["status"] == "awaiting_review"
    assert values["strategy"] == "evidence"
    assert values["rows"] == 4
    assert (values["rows_green"], values["rows_yellow"], values["rows_red"]) == (1, 1, 2)
    assert (values["reviewed_rejected"], values["review_pending"]) == (1, 1)
    assert values["coverage_oe"] == 0.75
    assert values["coverage_should_cost"] == 1.0
    assert (values["llm_calls"], values["llm_cost_usd"]) == (12, 0.0021)
    assert values["elapsed_s"] == 90.0
