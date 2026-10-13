"""Acceptance (spec 17 P9): the pages drive a real ``rfq api serve`` process in fake mode.

Upload 20 rows -> progress -> overview -> worksheet -> row detail -> review (including a new
OE number, finished by a worker process) -> download the worksheet.
"""

from __future__ import annotations

import io
import os
import socket
import subprocess
import sys
import time
from collections.abc import Callable, Iterator
from pathlib import Path

import httpx
import pytest
import yaml
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from dashboard.api_client import ApiClient, ApiError
from tests.dashboard.conftest import open_page, texts

WAIT_S = 120


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port: int = sock.getsockname()[1]
        return port


def _wait(check: Callable[[], bool], what: str) -> None:
    deadline = time.monotonic() + WAIT_S
    while not check():
        assert time.monotonic() < deadline, f"timed out waiting for {what}"
        time.sleep(0.3)


@pytest.fixture
def served_api(config_dir: Path, generated_dataset: Path, tmp_path: Path) -> Iterator[ApiClient]:
    """``rfq api serve`` in its own process, its workers in further processes."""
    path = config_dir / "settings.yaml"
    data = yaml.safe_load(path.read_text())
    data["data"]["dataset_dir"] = str(generated_dataset)
    path.write_text(yaml.safe_dump(data))
    port = _free_port()
    log = (tmp_path / "api.log").open("wb")
    server = subprocess.Popen(
        [sys.executable, "-m", "rfq_agent.cli", "api", "serve", "--port", str(port)],
        stdout=log,
        stderr=subprocess.STDOUT,
        env={**os.environ, "RFQ_CONFIG_DIR": str(config_dir)},
    )
    api = ApiClient(httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=30))

    def up() -> bool:
        assert server.poll() is None, (tmp_path / "api.log").read_text()
        try:
            api.list_runs()
        except ApiError:
            return False
        return True

    try:
        _wait(up, "the API server")
        yield api
    finally:
        server.terminate()
        server.wait(timeout=30)
        log.close()


def _poll_page(at: AppTest, done: Callable[[str], bool], what: str) -> None:
    """Rerun the page (as its polling fragment would) until its text satisfies ``done``."""

    def check() -> bool:
        at.run()
        return done(texts(at))

    _wait(check, what)


def test_the_whole_flow_on_the_pages(served_api: ApiClient, generated_dataset: Path) -> None:
    api = served_api
    sample = (generated_dataset / "rfq_sample_20.xlsx").read_bytes()

    # 1. upload and progress
    at = open_page("upload", api)
    at.file_uploader[0].upload("rfq_sample_20.xlsx", sample, "application/octet-stream").run()
    at.number_input(key="upload_batch_size").set_value(7).run()
    at.button(key="upload_start").click().run()
    run_id = at.session_state["run_id"]
    _poll_page(at, lambda shown: "已收尾" in shown, "the run to finalize")
    assert at.get("progress")[0].proto.value == 100

    # 2. overview
    at = open_page("overview", api, run_id=run_id)
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["总行数"] == "20"
    pending = int(metrics["待复核"])
    assert pending == int(metrics["红档"]) >= 2

    # 3. worksheet
    at = open_page("worksheet", api, run_id=run_id)
    assert list(at.dataframe[0].value["row_id"]) == list(range(1, 21))
    at.multiselect(key="ws_bands").set_value(["red"]).run()
    red = list(at.dataframe[0].value["row_id"])
    assert len(red) == pending

    # 4. row detail
    at = open_page("row", api, run_id=run_id, row_id=red[0])
    assert f"第 {red[0]} 行" in texts(at)
    assert at.get("vega_lite_chart")  # price, cost and signal charts

    # 5. review: accept one red row, give another a new OE number
    at = open_page("review", api, run_id=run_id)
    at.selectbox(key="rq_row").set_value(red[0]).run()
    at.radio(key="rq_action").set_value("accept").run()
    at.button(key="rq_submit").click().run()
    assert f"第 {red[0]} 行：已接受" in texts(at)

    at.selectbox(key="rq_row").set_value(red[1]).run()
    at.radio(key="rq_action").set_value("modify_oe").run()
    at.text_input(key="rq_oe").set_value("NEW-OE-1")
    at.button(key="rq_submit").click().run()
    assert "后台处理中" in texts(at)
    _poll_page(at, lambda shown: "改 OE 号已完成" in shown, "the new OE number")
    row = api.row(run_id, red[1])["row"]
    assert row["oe"]["current_oe"] == "NEW-OE-1"
    assert row["worksheet"]["review_status"] == "modified"

    # 6. evaluation: none before P10
    assert "暂无评估" in texts(open_page("evaluation", api, run_id=run_id))

    # 7. the worksheet download carries the reviews
    at = open_page("worksheet", api, run_id=run_id)
    assert [b.label for b in at.get("download_button")] == ["下载工作表（xlsx）"]
    sheet = load_workbook(io.BytesIO(api.worksheet(run_id)))["worksheet"]
    header = [c.value for c in sheet[1]]
    statuses = {
        r[header.index("row_id")]: r[header.index("review_status")]
        for r in sheet.iter_rows(min_row=2, values_only=True)
    }
    assert len(statuses) == 20
    assert (statuses[red[0]], statuses[red[1]]) == ("accepted", "modified")
