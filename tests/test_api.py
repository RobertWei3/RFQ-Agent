"""FastAPI endpoints (spec 15.1; phase-8 Q4, Q9, Q11, Q16-Q20), fake mode, in-process worker."""

from __future__ import annotations

import contextlib
import io
import os
import re
import signal
import threading
import time
from collections.abc import Callable, Iterator
from pathlib import Path

import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from rfq_agent.api import create_app
from rfq_agent.cli import app as cli_app
from rfq_agent.config import Settings, get_settings, load_settings
from rfq_agent.intake import read_rfq
from rfq_agent.schemas import RunDetail, RunEval, RunRecord, RunSpec
from rfq_agent.store.runs import RunStore

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ThreadLauncher:
    """Runs the worker's CLI command in a thread of this process; ``join`` waits for all."""

    def __init__(self) -> None:
        self.threads: list[threading.Thread] = []
        self.argvs: list[list[str]] = []
        self.failures: list[BaseException] = []

    def launch(self, argv: list[str], log_path: Path) -> None:
        self.argvs.append(argv)
        thread = threading.Thread(target=self._run, args=(argv,))
        self.threads.append(thread)
        thread.start()

    def _run(self, argv: list[str]) -> None:
        try:
            # A worker has no terminal: any prompt must fail, not wait.
            with contextlib.redirect_stdout(io.StringIO()):
                cli_app(argv, standalone_mode=False)
        except BaseException as exc:  # noqa: BLE001 - reported by join()
            self.failures.append(exc)

    def join(self) -> None:
        for thread in self.threads:
            thread.join(timeout=120)
        assert self.failures == []


class NullLauncher:
    """Starts nothing: the worker never runs (as if it died before taking over)."""

    def __init__(self) -> None:
        self.argvs: list[list[str]] = []

    def launch(self, argv: list[str], log_path: Path) -> None:
        self.argvs.append(argv)


@pytest.fixture
def api_settings(config_dir: Path, generated_dataset: Path) -> Settings:
    path = config_dir / "settings.yaml"
    data = yaml.safe_load(path.read_text())
    data["data"]["dataset_dir"] = str(generated_dataset)
    path.write_text(yaml.safe_dump(data))
    get_settings.cache_clear()
    return load_settings(config_dir)


@pytest.fixture
def launcher() -> ThreadLauncher:
    return ThreadLauncher()


@pytest.fixture
def client(api_settings: Settings, launcher: ThreadLauncher) -> Iterator[TestClient]:
    with TestClient(create_app(api_settings, launcher)) as c:
        yield c


def _sample(generated_dataset: Path) -> bytes:
    return (generated_dataset / "rfq_sample_20.xlsx").read_bytes()


def _post_run(client: TestClient, generated_dataset: Path, **form: str) -> RunRecord:
    files = {"file": ("rfq_sample_20.xlsx", _sample(generated_dataset), XLSX)}
    response = client.post("/runs", files=files, data={"batch_size": "7", **form})
    assert response.status_code == 202, response.text
    return RunRecord.model_validate(response.json())


def _finished_run(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path, **form: str
) -> RunDetail:
    record = _post_run(client, generated_dataset, **form)
    launcher.join()
    return RunDetail.model_validate(client.get(f"/runs/{record.run_id}").json())


def _store(settings: Settings) -> RunStore:
    return RunStore.open(settings.db_path)


# --- runs -----------------------------------------------------------------------------------


def test_posting_an_rfq_starts_a_background_run_that_finishes(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    record = _post_run(client, generated_dataset)

    assert record.status == "running"
    assert record.spec.batch_size == 7 and record.spec.review is False
    assert Path(record.spec.input).read_bytes() == _sample(generated_dataset)
    assert launcher.argvs[0][:3] == ["run", "--resume", record.run_id]
    launcher.join()

    detail = RunDetail.model_validate(client.get(f"/runs/{record.run_id}").json())
    assert detail.run.status == "completed"
    assert (detail.progress.batches_total, detail.progress.completed) == (3, 3)
    assert detail.summary is not None and detail.summary.rows == 20
    assert sum(detail.summary.bands.values()) == 20
    assert detail.lease_active is False
    listed = [RunRecord.model_validate(r) for r in client.get("/runs").json()]
    assert [r.run_id for r in listed] == [record.run_id]


def test_a_bad_upload_is_rejected_and_creates_no_run(
    client: TestClient, launcher: ThreadLauncher
) -> None:
    not_excel = client.post("/runs", files={"file": ("rfq.txt", b"hello", "text/plain")})
    book = Workbook()
    assert book.active is not None
    book.active.append(["part", "qty"])
    buffer = io.BytesIO()
    book.save(buffer)
    no_column = client.post("/runs", files={"file": ("rfq.xlsx", buffer.getvalue(), XLSX)})
    bad_strategy = client.post(
        "/runs", files={"file": ("rfq.xlsx", buffer.getvalue(), XLSX)}, data={"strategy": "x"}
    )

    assert (not_excel.status_code, no_column.status_code, bad_strategy.status_code) == (
        422,
        422,
        422,
    )
    assert "customer_part_no" in no_column.json()["detail"]
    assert client.get("/runs").json() == []
    assert launcher.argvs == []


def test_a_real_provider_needs_the_server_started_with_allow_real(
    real_provider_env: None, client: TestClient, generated_dataset: Path
) -> None:
    files = {"file": ("rfq_sample_20.xlsx", _sample(generated_dataset), XLSX)}

    response = client.post("/runs", files=files)

    assert response.status_code == 409
    assert "--allow-real" in response.json()["detail"]
    assert "oe_interchange" in response.json()["detail"]
    assert client.get("/runs").json() == []


def test_unknown_runs_and_rows_are_not_found(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    detail = _finished_run(client, launcher, generated_dataset)
    run_id = detail.run.run_id

    for path in (
        "/runs/nope",
        "/runs/nope/rows",
        "/runs/nope/worksheet.xlsx",
        f"/runs/{run_id}/rows/999",
        f"/runs/{run_id}/eval",  # not evaluated yet
        "/runs/nope/eval",
    ):
        assert client.get(path).status_code == 404, path
    assert client.post("/runs/nope/resume").status_code == 404
    review = client.post(f"/runs/{run_id}/rows/999/review", json={"action": "accept"})
    assert review.status_code == 404


def test_eval_returns_the_stored_evaluation_once_the_run_is_evaluated(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path, tmp_path: Path
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset).run.run_id
    assert client.get(f"/runs/{run_id}/eval").status_code == 404

    done = CliRunner().invoke(cli_app, ["eval", "--run", run_id, "--out", str(tmp_path)])
    assert done.exit_code == 0, done.output

    response = client.get(f"/runs/{run_id}/eval")
    assert response.status_code == 200
    body = RunEval.model_validate(response.json())
    assert body.run_id == run_id and body.rows == 20
    assert body.truth_path.endswith("rfq_sample_20_ground_truth.jsonl")


# --- rows and worksheet -----------------------------------------------------------------------


def test_rows_filter_by_band_and_review_status_and_page_in_row_order(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset).run.run_id

    red = client.get(f"/runs/{run_id}/rows", params={"band": "red", "status": "unreviewed"}).json()
    page = client.get(f"/runs/{run_id}/rows", params={"offset": 5, "limit": 5}).json()
    both = client.get(f"/runs/{run_id}/rows", params=[("band", "red"), ("band", "green")]).json()

    assert red["total"] > 0 and all(r["band"] == "red" for r in red["items"])
    assert (page["total"], page["offset"], page["limit"]) == (20, 5, 5)
    assert [r["row_id"] for r in page["items"]] == [6, 7, 8, 9, 10]
    assert {r["band"] for r in both["items"]} <= {"red", "green"}
    assert client.get(f"/runs/{run_id}/rows", params={"band": "blue"}).status_code == 422
    assert client.get(f"/runs/{run_id}/rows", params={"limit": 1001}).status_code == 422


def test_a_row_shows_its_full_detail(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset).run.run_id

    detail = client.get(f"/runs/{run_id}/rows/3").json()

    assert detail["row"]["line"]["row_id"] == 3
    assert detail["row"]["confidence"]["signals"]
    assert detail["row"]["oe"]["reasoning"]
    assert (detail["reviews"], detail["in_review_queue"]) == ([], False)


def test_the_worksheet_downloads_once_the_run_is_finalized(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path, api_settings: Settings
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset).run.run_id
    store = _store(api_settings)
    pending = store.create_run(
        RunSpec(input="x.xlsx", dataset_dir="x", batch_size=7, strategy="evidence", review=False),
        read_rfq(generated_dataset / "rfq_sample_20.xlsx").lines,
    )
    store.close()

    response = client.get(f"/runs/{run_id}/worksheet.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert f"run_{run_id}.xlsx" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert [s.title for s in book.worksheets] == ["worksheet", "summary"]
    assert book.worksheets[0].max_row == 21
    assert client.get(f"/runs/{pending}/worksheet.xlsx").status_code == 409


# --- review -----------------------------------------------------------------------------------


def _queue(client: TestClient, run_id: str) -> list[int]:
    rows = client.get(f"/runs/{run_id}/rows", params={"band": "red", "status": "unreviewed"})
    return [r["row_id"] for r in rows.json()["items"]]


def test_any_queued_row_can_be_reviewed_through_the_api(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    detail = _finished_run(client, launcher, generated_dataset, review="true")
    run_id = detail.run.run_id
    queue = _queue(client, run_id)
    assert detail.run.status == "awaiting_review" and len(queue) >= 3
    assert detail.summary is not None and detail.summary.review_pending == len(queue)
    assert client.get(f"/runs/{run_id}/rows/{queue[-1]}").json()["in_review_queue"] is True

    rejected = client.post(
        f"/runs/{run_id}/rows/{queue[-1]}/review", json={"action": "reject", "note": "not ours"}
    )
    priced = client.post(
        f"/runs/{run_id}/rows/{queue[0]}/review", json={"action": "modify_price", "msrp": 42.0}
    )

    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["done"] is True
    assert rejected.json()["row"]["worksheet"]["review_status"] == "rejected"
    assert priced.json()["row"]["worksheet"]["msrp"] == 42.0
    again = client.post(f"/runs/{run_id}/rows/{queue[-1]}/review", json={"action": "accept"})
    assert again.status_code == 409
    green = client.get(f"/runs/{run_id}/rows", params={"band": "green"}).json()["items"][0]
    not_red = client.post(
        f"/runs/{run_id}/rows/{green['row_id']}/review", json={"action": "accept"}
    )
    assert not_red.status_code == 409
    incomplete = client.post(
        f"/runs/{run_id}/rows/{queue[1]}/review", json={"action": "modify_price"}
    )
    assert incomplete.status_code == 422
    entries = client.get(f"/runs/{run_id}/rows/{queue[-1]}").json()["reviews"]
    assert [e["action"] for e in entries] == ["reject"]

    for rid in _queue(client, run_id):
        assert client.post(
            f"/runs/{run_id}/rows/{rid}/review", json={"action": "accept"}
        ).is_success
    assert client.get(f"/runs/{run_id}").json()["run"]["status"] == "completed"


def test_a_new_oe_number_is_accepted_and_finished_by_a_worker(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset, review="true").run.run_id
    rid = _queue(client, run_id)[0]

    response = client.post(
        f"/runs/{run_id}/rows/{rid}/review", json={"action": "modify_oe", "oe": "manual-oe-1"}
    )

    assert response.status_code == 202, response.text
    assert (response.json()["done"], response.json()["row"]) == (False, None)
    assert launcher.argvs[-1][:3] == ["runs", "answer", run_id]
    launcher.join()
    row = client.get(f"/runs/{run_id}/rows/{rid}").json()["row"]
    assert row["worksheet"]["review_status"] == "modified"
    assert (row["worksheet"]["oe_status"], row["worksheet"]["current_oe"]) == (
        "resolved",
        "MANUAL-OE-1",
    )


def test_review_is_refused_without_review_mode(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset).run.run_id
    red = client.get(f"/runs/{run_id}/rows", params={"band": "red"}).json()["items"][0]

    response = client.post(f"/runs/{run_id}/rows/{red['row_id']}/review", json={"action": "accept"})

    assert response.status_code == 409


# --- run lease and resume ---------------------------------------------------------------------


def test_a_run_another_process_drives_cannot_be_reviewed_or_resumed(
    client: TestClient, launcher: ThreadLauncher, generated_dataset: Path, api_settings: Settings
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset, review="true").run.run_id
    rid = _queue(client, run_id)[0]
    store = _store(api_settings)
    store.acquire_lease(run_id, "cli", ttl_s=api_settings.storage.lease_ttl_s)

    review = client.post(f"/runs/{run_id}/rows/{rid}/review", json={"action": "accept"})
    detail = client.get(f"/runs/{run_id}").json()

    assert review.status_code == 409
    assert "is driven by cli" in review.json()["detail"]
    assert detail["lease_active"] is True
    store.release_lease(run_id, "cli")
    store.close()
    assert client.post(f"/runs/{run_id}/rows/{rid}/review", json={"action": "accept"}).is_success


def test_resume_restarts_a_run_whose_worker_is_gone(
    api_settings: Settings, generated_dataset: Path
) -> None:
    store = _store(api_settings)
    run_id = store.create_run(
        RunSpec(
            input=str(generated_dataset / "rfq_sample_20.xlsx"),
            dataset_dir=str(generated_dataset),
            batch_size=7,
            strategy="evidence",
            review=False,
        ),
        read_rfq(generated_dataset / "rfq_sample_20.xlsx").lines,
    )
    store.close()
    launcher = ThreadLauncher()
    with TestClient(create_app(api_settings, launcher)) as client:
        resumed = client.post(f"/runs/{run_id}/resume")
        assert resumed.status_code == 202, resumed.text
        launcher.join()

        assert client.get(f"/runs/{run_id}").json()["run"]["status"] == "completed"
        assert client.post(f"/runs/{run_id}/resume").status_code == 409  # nothing to resume


def test_a_run_whose_worker_has_not_taken_over_yet_is_not_resumed_twice(
    api_settings: Settings, generated_dataset: Path
) -> None:
    with TestClient(create_app(api_settings, NullLauncher())) as client:
        record = _post_run(client, generated_dataset)

        response = client.post(f"/runs/{record.run_id}/resume")

        assert response.status_code == 409
        assert client.get(f"/runs/{record.run_id}").json()["lease_active"] is True


def test_the_openapi_document_lists_every_endpoint(client: TestClient) -> None:
    paths = client.get("/openapi.json").json()["paths"]

    assert set(paths) == {
        "/runs",
        "/runs/{run_id}",
        "/runs/{run_id}/rows",
        "/runs/{run_id}/rows/{row_id}",
        "/runs/{run_id}/rows/{row_id}/review",
        "/runs/{run_id}/worksheet.xlsx",
        "/runs/{run_id}/eval",
        "/runs/{run_id}/resume",
    }
    assert paths["/runs/{run_id}"]["get"]["responses"]["200"]["content"]["application/json"][
        "schema"
    ]["$ref"].endswith("/RunDetail")
    assert paths["/runs/{run_id}/eval"]["get"]["responses"]["200"]["content"]["application/json"][
        "schema"
    ]["$ref"].endswith("/RunEval")


# --- acceptance: real worker processes (spec 13 / phase-8 seam 6) -----------------------------


def _wait_for(check: Callable[[], bool], what: str, seconds: float = 120) -> None:
    deadline = time.monotonic() + seconds
    while not check():
        assert time.monotonic() < deadline, f"timed out waiting for {what}"
        time.sleep(0.05)


def _worker_pid(settings: Settings, run_id: str) -> int:
    store = _store(settings)
    try:
        lease = store.lease(run_id)
    finally:
        store.close()
    assert lease is not None
    return int(lease.owner.split("@")[0])


def test_api_workers_run_in_the_background_and_a_killed_one_resumes(
    api_settings: Settings, generated_dataset: Path, tmp_path: Path
) -> None:
    """Spec 13 kill/resume through the API: same rows as an unbroken CLI run."""
    rfq = generated_dataset / "rfq_200.xlsx"
    out = tmp_path / "cli-out"
    args = ["run", "--input", str(rfq), "--batch-size", "20", "--output-dir", str(out)]
    unbroken = CliRunner().invoke(cli_app, args)
    assert unbroken.exit_code == 0, unbroken.output
    found = re.search(r"run_id=(\S+)", unbroken.output)
    assert found, unbroken.output
    reference = (out / f"run_{found.group(1)}.jsonl").read_text()

    with TestClient(create_app(api_settings)) as client:
        files = {"file": ("rfq_200.xlsx", rfq.read_bytes(), XLSX)}
        posted = client.post("/runs", files=files, data={"batch_size": "20"})
        run_id = posted.json()["run_id"]

        def completed() -> int:
            return int(client.get(f"/runs/{run_id}").json()["progress"]["completed"])

        _wait_for(lambda: completed() >= 3, "three batches")
        pid = _worker_pid(api_settings, run_id)
        assert pid != os.getpid()
        refused = CliRunner().invoke(cli_app, ["run", "--resume", run_id])
        assert refused.exit_code == 3, refused.output
        assert "is driven by" in refused.output

        os.kill(pid, signal.SIGKILL)
        _wait_for(lambda: not _alive(pid), "the worker to die")
        detail = client.get(f"/runs/{run_id}").json()
        assert detail["run"]["status"] == "running"
        assert detail["lease_active"] is False
        done_before = detail["progress"]["completed"]
        assert 3 <= done_before < 10

        assert client.post(f"/runs/{run_id}/resume").status_code == 202
        _wait_for(
            lambda: client.get(f"/runs/{run_id}").json()["run"]["status"] == "completed",
            "the resumed run",
        )

        store = _store(api_settings)
        batches = store.batches(run_id)
        store.close()
        assert sum(b.attempts == 1 for b in batches) >= done_before
        exported = api_settings.db_path.parent / "runs" / f"run_{run_id}.jsonl"
        # The worker exports after the run is marked completed; the meta file comes last.
        _wait_for(exported.with_suffix(".meta.json").exists, "the worker's export")
        assert exported.read_text() == reference
        log = api_settings.db_path.parent / "runs" / f"run_{run_id}.worker.log"
        assert "batches: completed 10" in log.read_text()


def _alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    return True


def _real_settings(monkeypatch: pytest.MonkeyPatch, config_dir: Path) -> Settings:
    monkeypatch.delenv("RFQ_LLM_PROVIDER", raising=False)
    get_settings.cache_clear()
    return load_settings(config_dir)


def test_work_that_calls_a_real_provider_is_refused_with_its_expected_calls(
    client: TestClient,
    launcher: ThreadLauncher,
    generated_dataset: Path,
    api_settings: Settings,
    config_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    run_id = _finished_run(client, launcher, generated_dataset, review="true").run.run_id
    rid = _queue(client, run_id)[0]
    store = _store(api_settings)
    stopped = store.create_run(
        RunSpec(input="x.xlsx", dataset_dir="x", batch_size=7, strategy="evidence", review=False),
        read_rfq(generated_dataset / "rfq_sample_20.xlsx").lines,
    )
    store.close()
    real = _real_settings(monkeypatch, config_dir)

    with TestClient(create_app(real, launcher)) as paid:
        new_oe = paid.post(
            f"/runs/{run_id}/rows/{rid}/review", json={"action": "modify_oe", "oe": "X-1"}
        )
        resumed = paid.post(f"/runs/{stopped}/resume")

    assert (new_oe.status_code, resumed.status_code) == (409, 409)
    assert "oe_msrp ~1" in new_oe.json()["detail"]
    assert "--allow-real" in new_oe.json()["detail"]
    assert "oe_interchange ~60" in resumed.json()["detail"]
    assert len(launcher.argvs) == 1  # only the fake run's own worker
